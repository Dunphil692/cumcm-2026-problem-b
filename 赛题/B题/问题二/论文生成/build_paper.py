"""按 2026 国赛标准论文模板生成问题二论文 docx。"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "pylibs"))

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
EQ = os.path.join(HERE, "eq")
FIGS = os.path.join(ROOT, "figs")
CODE = os.path.join(ROOT, "code")

doc = Document()

# ---------------- 页面 ----------------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.54)
sec.left_margin = sec.right_margin = Cm(3.17)

# ---------------- 样式 ----------------
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(0)


def set_run(run, size=12, bold=False, east="宋体", latin="Times New Roman", italic=False):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)


def para(text="", size=12, bold=False, align=None, indent=True, east="宋体",
         latin="Times New Roman", spacing=1.5, before=0, after=0):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = spacing
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, east=east, latin=latin)
    return p


def rich_para(parts, indent=True, align=None, spacing=1.5):
    """parts: list of (text, dict(bold/size/east/latin/italic))"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = spacing
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    if align is not None:
        p.alignment = align
    for text, kw in parts:
        r = p.add_run(text)
        set_run(r, size=kw.get("size", 12), bold=kw.get("bold", False),
                east=kw.get("east", "宋体"), latin=kw.get("latin", "Times New Roman"),
                italic=kw.get("italic", False))
    return p


def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    set_run(r, size=15, bold=True, east="黑体", latin="Times New Roman")
    return p


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    set_run(r, size=13, bold=True, east="黑体", latin="Times New Roman")
    return p


def caption(text, before=6, after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run(r, size=10.5, bold=True, east="黑体", latin="Times New Roman")
    return p


def figure(path, cap, width_cm=14.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    if cap is not None:
        caption(cap, before=2, after=8)
    else:
        p.paragraph_format.space_after = Pt(8)


def equation(png, number, width_cm=None):
    from PIL import Image
    if width_cm is None:
        im = Image.open(png)
        w_px, h_px = im.size
        width_cm = min(14.5, w_px / 300.0 * 2.54)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    # 右侧制表位放编号
    p.paragraph_format.tab_stops.add_tab_stop(Cm(14.6), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run()
    run.add_picture(png, width=Cm(width_cm))
    r = p.add_run("\t(%s)" % number)
    set_run(r, size=12)
    return p


def three_line_table(headers, rows, widths=None, fontsize=10.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    data = [headers] + rows
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(val))
            set_run(r, size=fontsize, bold=(i == 0), east="宋体", latin="Times New Roman")
    # 三线表边框
    tbl = t._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag, sz in (("top", 12), ("bottom", 12)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for tag in ("left", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)
    # 表头下边线
    for cell in t.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        el = OxmlElement("w:bottom")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "000000")
        b.append(el)
        tcPr.append(b)
    return t


def add_page_number_footer():
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run_el = OxmlElement("w:r")
    t_el = OxmlElement("w:t")
    t_el.text = "1"
    run_el.append(t_el)
    fld.append(run_el)
    p._p.append(fld)


add_page_number_footer()

# ================= 第一页：题目 + 摘要 =================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(10)
r = p.add_run("基于保证接收域与最坏交会直径的第二检测点选择策略研究")
set_run(r, size=16, bold=True, east="黑体")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("摘  要")
set_run(r, size=15, bold=True, east="黑体")

para("无线电干扰源的自动定位与清除任务中，检测设备在单个检测点只能获得干扰源的示向度（信号到达方向，测量误差 ±1°）而无法测距，且仅当干扰源处于有效接收半径（1000~1500 m，具体值未知）内时才能收到信号。单次观测只能确定干扰源所在的一条角度带，无法确定其位置。本文研究第二检测点的选择问题：给定第一个检测点 S1 及其示向度 θ1，如何选择第二检测点 S2，使两站交会定位效果“较好”（定位区域直径尽可能小），并给出 S2 的候选区域。",
     size=12, indent=True)

para("首先，综合示向楔形、接收半径上下界与目标区域边界，建立干扰源的源可能域 Ω1，并对目标圆进行逐方向裁剪（切向构型下 89°/90°/91° 方向的退出距离分别为 160.79 / 189.47 / 223.27 m）。其次，利用“第一检测点已收到信号 ⟹ 接收半径不小于源距”这一关键信息，推导第二检测点的一般保证接收域 C_recv：对源可能域内每个源 g 须满足 d2 ≤ max(1000, ‖g−S1‖)，并证明常用三圆交区域是该域的保守充分条件。然后，以所有可能源、所有相容测向误差下的定位区域直径最大值为目标函数 J（定位区域严格裁剪回源可能域与接收上界，排除物理上不可能的位置），在保证接收域内通过多级网格与细化进行确定性寻优。数值结果表明：中心构型（S1 位于目标区域中心、示向沿正东）下最优第二检测点为 (844, ±544) m，偏角 32.8°、离 S1 约 1004 m，最坏定位区域直径 Jmin = 111.0 m；向外 800 m 构型与切向边界构型下 Jmin 分别为 41.0 m 与 8.2 m。以 J ≤ 1.25 Jmin 的近最优水平集给出候选区域：中心构型为离 S1 约 855~1004 m、偏角 16.1°~60.1° 的两瓣，敏感性分析表明该区域形状对容差取值（1.1~1.5）不敏感。",
     size=12, indent=True)

para("本文模型将接收保证、区域裁剪与最坏情形口径统一为可由几何精确计算的确定性框架：无需随机搜索即可稳定求解，结果完全可复现；近场强信号、逐方向目标圆裁剪、退化交会等边界情形处理完备。模型检验表明：与另一套独立实现（不同几何引擎与采样方案）的三场景交叉验证中，最坏直径相差不超过 0.006 m；最坏情形一致出现在楔形远弧端点与误差极值处；容差灵敏度与边界回退稳健性均通过。模型可直接推广到一般的第一检测点位置，为问题三的多点序列定位与清除提供粗定位基础。",
     size=12, indent=True)

p = rich_para([("关键词：", {"bold": True, "east": "黑体"}),
               ("示向交会定位；保证接收域；最坏定位区域直径；近最优水平集；敏感性分析", {})],
              indent=False, align=None)
p.paragraph_format.space_before = Pt(10)

doc.add_page_break()

# ================= 问题重述 =================
h1("问题重述")
para("为清除目标区域内的无线电干扰源，机器狗携带的检测设备需要在不同检测点接收干扰源信号并估计其位置。设备在检测点只能测得干扰源的示向度（信号到达方向），测量存在 ±1° 的误差；仅当干扰源处于设备的有效接收半径内时才能收到信号，该半径介于 1000 m 与 1500 m 之间但具体数值未知；干扰源与检测点的距离不超过 5 m 时信号过强、无法获得示向度。目标区域是以原点为圆心、半径 1800 m 的圆。")
para("已知第一个检测点 S1 及其示向度 θ1。单次观测只给出方向、不给出距离，因此只能确定干扰源位于沿示向线张角 2° 的一条楔形区域内。为提高交会定位的精度，需要选择第二个检测点 S2，通过两站示向线交会缩小干扰源的可能范围。")
para("本文要解决的问题是：（1）给出第二检测点 S2 的选择策略，使两站交会定位的定位区域直径尽可能小，同时保证 S2 能收到干扰源的信号；（2）给出 S2 的候选区域，为后续多点序列定位与清除提供基础。其中定位区域直径沿用问题一的定义：定位区域内任意两点之间距离的最大值。")

# ================= 1 问题分析 =================
h1("1  问题分析")
h2("1.1  问题二的分析")
para("本问题的实质是在三个约束下求一个最坏情形的最小化。约束一“听得见”：S2 必须保证能收到干扰源信号——由于接收半径未知且源位置未知，必须对源可能域内的每个位置逐一保证，这是本问题最容易被简化错的地方（常用三圆近似只覆盖了部分源）。约束二“交会角不能太小”：两次示向线若近平行，交会区域被拉成细长条、直径巨大。约束三（可选）“留在目标区域内”：减少机器狗移动、留在搜索区内。目标“较好”的含义：任务要求稳健地缩小源位置范围且接收半径未知，因此以所有可能情形下的最坏定位区域直径作为评价指标（最坏口径），并在其上求最小。")
para("建模思路分四步：先刻画源可能域 Ω1（源可能在哪里），再刻画保证接收域 C_recv（S2 在哪些位置能保证收到信号），再定义最坏直径目标 J（定位效果有多好），最后确定性寻优并给出候选区域（近最优水平集）。三个关键点需要强调：其一，保证接收的推导——S1 收到信号意味着接收半径 R ≥ ‖G−S1‖，因此对远源 S2 的距离约束可放宽到 1500 m，而对近源必须保持 1000 m 以内，“一律按 1000 m”过保守、“一律按 1500 m”是假保证；其二，定位区域必须裁剪回源可能域与接收上界——两楔形交会多边形可能延伸到源不可能存在的位置（离 S1 超过 1500 m），直接计算其直径会严重高估；其三，候选区域——严格最优只有离散的点，题目要求给出“区域”，故采用近最优水平集并明确容差的含义。整体思路框架见图 1。")
figure(os.path.join(HERE, "flowchart.png"), None, width_cm=11.0)  # 图中已含“图 1”标题

# ================= 2 模型假设 =================
h1("2  模型假设")
assumptions = [
    "示向度测量误差不超过 ±1°（题面给定）；两次观测误差均在此界内，并按有界最坏情形处理，不假设其概率分布。",
    "有效接收半径是区间 [1000, 1500] m 内的一个固定未知值；第一检测点已收到源信号，故该半径不小于源与 S1 的距离。",
    "干扰源位于半径 1800 m 的目标圆内，且与 S1 的距离大于 5 m（否则无法获得示向度）。",
    "评价口径采用所有相容情形下的最坏值，与“确保清除”的任务目标一致。",
    "建模自加假设（题面未强制）：S2 不进入第一条示向楔形（偏角大于 1°，避免两次示向近平行导致交会退化）；S2 留在目标圆内（减少移动、留在搜索区内；三个构型中该约束均非紧）。",
    "定位区域直径沿用问题一的定义；本文以直径（而非面积）作为评价指标——面积无法区分细长条与团块，直径直接刻画最坏定位精度。",
]
for i, a in enumerate(assumptions, 1):
    rich_para([(f"（{i}）", {"bold": True}), (a, {})])

# ================= 3 符号说明 =================
h1("3  符号说明")
caption("表 1  本文的符号说明", before=6, after=3)
sym_rows = [
    ("Ω1", "干扰源源可能域", "—"),
    ("g / G", "源可能域内一个可能的源位置 / 真实干扰源", "—"),
    ("S1, S2", "第一、第二检测点", "—"),
    ("θ1, θ2", "第一、第二次示向度", "°"),
    ("d1, d2", "源到 S1、S2 的距离", "m"),
    ("R", "有效接收半径（未知，1000~1500）", "m"),
    ("L(g)", "保证接收半径，L(g)=max(1000, d1)", "m"),
    ("C_recv", "一般保证接收域", "—"),
    ("F(q)", "接收约束函数 sup(‖q−g‖−L(g))", "m"),
    ("W1, W2", "两站示向楔形（各 ±1°）", "—"),
    ("Ω2", "第二次观测后的定位区域", "—"),
    ("J(q) / Jmin", "最坏定位区域直径 / 其最小值", "m"),
    ("e2", "第二次示向度误差", "°"),
    ("φ", "源方向相对示向线的偏角", "°"),
    ("r_far", "楔形逐方向的远端距离", "m"),
    ("α", "候选区域容差系数", "1"),
    ("ψ", "交会角（G 处两示向线夹角）", "°"),
]
three_line_table(["符号", "说明", "单位"], sym_rows)
para("")

# ================= 4 模型建立与求解 =================
h1("4  模型建立与求解")

h2("4.1  源可能域 Ω1 的建立")
para("第一次观测只给出示向度：源必位于沿示向线、张角 2° 的楔形内；S1 收到信号要求 5 < d1 ≤ R ≤ 1500；源又在目标圆内。故源可能域为")
equation(os.path.join(EQ, "eq1.png"), 1)
para("注意目标圆以原点为心、不以 S1 为心：楔形内各方向被目标圆截断的距离不同，必须逐方向计算该方向的远端距离 r_far(φ)=min(1500, 该方向穿出目标圆的距离)。本文考察三个典型构型：中心构型（S1=(0,0)，θ1=0°）各方向远端均为 1500 m；向外构型（S1=(−1000,0)，θ1=180°）楔形指向圆外，中线方向在 800 m 处穿出目标圆、±1° 两侧约 800.07 m，不能统一按 1500 m 截断；切向构型（S1=(1790,0)，θ1=90°）楔形贴近圆边，89°/90°/91° 三个方向的退出距离分别为 160.79 / 189.47 / 223.27 m，各不相同。上述解析退出距离与数值射线计算核对一致（最大差异 6.3×10⁻¹³ m）。")

h2("4.2  保证接收域 C_recv 的建立")
para("S2 的“听得见”必须对 Ω1 内每个可能的源 g 都成立。由 S1 已收到信号知 R ≥ d1，且 R ≤ 1500；要使 S2 对一切相容的 R 都能收到 g，必须且只需 d2 ≤ max(1000, d1)：")
equation(os.path.join(EQ, "eq2.png"), 2)
para("由此定义接收约束函数 F 与一般保证接收域 C_recv：")
equation(os.path.join(EQ, "eq3.png"), 3)
para("几何解释：对近源段（d1≤1000 m）约束为 d2≤1000，最紧的源是 S1 附近与 1000 m 远弧两端 P−、P+，即三圆叶子 B(S1,1000)∩B(P−,1000)∩B(P+,1000)——它是 C_recv 的保守充分条件：")
equation(os.path.join(EQ, "eq8.png"), 8)
para("三圆叶子对近源段精确（近源段 ⊂ 三锚点凸包，凸性保证盘约束成立），对远源段因 L(g)=d1 自动满足。C_recv 的真实外沿由近端 5 m 处的源逐点决定，可比 1000 m 圆盘略外扩约 4 m（中心构型最优点离 S1 1004 m）。数值实现中，以源可能域有限网格上 F 的上确界估计 F̂(q)≤0 作为判定，并如实标注其为数值口径而非连续认证。")

h2("4.3  目标函数：最坏定位区域直径")
para("第二次观测（非近场、非重复地点）后，干扰源的定位区域为两楔形之交再裁剪回源可能域与接收上界：")
equation(os.path.join(EQ, "eq4.png"), 4)
para("其中裁剪回 Ω1 排除了 d1>1500 m 的不可能位置，裁剪 d2≤1500 排除了 S2 收不到的位置。若某个可能的源 g 与 S2 的距离不超过 5 m，第二次观测为近场强信号（无示向度），定位区域取 Ω1∩B(q,5) 的实际交集。目标取所有可能源、所有相容第二次误差下的最坏情形：")
equation(os.path.join(EQ, "eq5.png"), 5)
para("小角度机理：两楔形近似为带宽 wi=2ri·tan1°、交角 ψ 的平行四边形，直径近似为")
equation(os.path.join(EQ, "eq7.png"), 7)
para("ψ→0°/180° 时 D→∞，ψ 接近 90° 且 r1、r2 不大时 D 小。听得见约束把 S2 锁在离 S1 约 1000 m 处，因此最优点落在“贴外沿 + 大偏角（交会角约 40°）”的折中位置。注意偏角（从 S1 看）与交会角（在 G 处两示向线夹角）不是同一个角。")
para("后验裁剪的必要性（数值实例）：若不裁剪，中心构型 S2=(800,600) 处的最坏情形（源在 1500 m、+1° 方向，第二次误差 +1°）的后验四顶点中有三个离 S1 达 1551~1623 m——物理上不可能（S1 收到源要求源距不超过 1500 m）——未裁剪直径被虚高至 134.98 m；裁剪后该情形退化为真源附近的单点，同一点的最坏直径降至 112.1 m。可见式(4) 中的域裁剪不是可有可无的细节，而是模型正确性的组成部分。")

h2("4.4  模型求解")
para("数值口径：源可能域按方向 φ∈[−1°,1°] 与每方向 5 m 到 r_far 的径向网格离散，第二次误差 e2∈[−1°,1°] 按网格离散，J 取该有限网格上的最大值（有限数值上确界估计，非连续认证上界）。与蒙特卡洛数值积分[3] 的随机误差不同，本文采用确定性网格，结果逐位可复现。求解分两步：先用 F̂(q)≤0、偏角 >1°、（可选）目标圆内三条规则做可行性预筛，再在可行域上做 25 m 粗网格 → 5 m → 1 m → 0.5 m 的多级细化寻优。由于目标是一维标量且可由几何精确计算，无需差分进化[1]、NSGA-II[2] 等随机多目标优化即可稳定求解，避免了随机搜索的复现性问题。")
para("三个构型的求解结果见表 2，候选区域见图 2~图 4。最坏情形见证（三构型一致）：最坏直径出现在楔形远弧端点（φ=±1°、d1=r_far）× 第二次误差取极值（e2=±1°）处。中心构型的最优点 (844,±544) 恰在 C_recv 边界上（数值裕量 0.02 m）：这是“基线尽量长但受接收约束”的自然结果。")
caption("表 2  三构型求解结果", before=8, after=3)
res_rows = [
    ("中心 (0,0), 0°", "111.0", "(844, ±544)", "1004 / 32.8°", "−0.02（贴边界）", "855~1004 / 16.1~60.1"),
    ("向外 800 (−1000,0), 180°", "41.0", "(−1594.8, ±511.4)", "784 / 40.7°", "−219.3", "593~1000 / 24.4~60.8"),
    ("切向 (1790,0), 90°", "8.2", "(1688.4, +120.5)", "158 / 40.1°", "−846.1", "125~236 / 31.0~60.3（西侧一瓣）"),
]
three_line_table(["构型（S1, θ1）", "Jmin (m)", "最优 S2 (m)", "离 S1 距离 (m) / 偏角", "接收裕量 (m)", "1.25 候选区：离 S1 (m) / 偏角"], res_rows)
para("")
para("三个构型的几何含义一致：最优 S2 均“贴保证接收域外沿、偏角 32°~41°”，在基线长度与交会角之间取得折中；向外构型因楔形被目标圆截短，源域变小、最坏直径显著下降；切向构型源域最窄，最坏直径仅 8.2 m。中心与向外构型关于示向线两侧对称；切向构型楔形轴的镜像点在目标圆外（不可用），候选区只在楔形西侧一瓣。")
figure(os.path.join(FIGS, "problem2_candidate_region_center_reference.png"), "图 2  中心构型候选区域（彩色为可行域最坏直径，等值线为 1.1/1.25/1.5×Jmin，★为最优点）", width_cm=12.5)
figure(os.path.join(FIGS, "problem2_candidate_region_inside_outward_800.png"), "图 3  向外 800 m 构型候选区域", width_cm=12.5)
figure(os.path.join(FIGS, "problem2_candidate_region_tangent_boundary.png"), "图 4  切向边界构型候选区域（西侧一瓣）", width_cm=12.5)

h2("4.5  候选区域：近最优水平集")
para("严格最优只有左右各一个点，单点交不了“候选区域”这道题，故把“较好”放宽为近最优水平集：")
equation(os.path.join(EQ, "eq6.png"), 6)
para("α=1.25 的含义是允许最坏直径比最优值差 25%（中心构型 111→139 m），属于建模容差，不是题面常数。敏感性分析（表 3）表明：α 在 1.1~1.5 之间时候选区域始终是贴外沿的两瓣，只改变向内伸入的深度——策略结论不依赖 α 的精确取值。这一结果与测量学前方交会“交会角 30°~150°”的经验准则给出一致的区域。")
caption("表 3  容差系数 α 的敏感性（中心构型）", before=8, after=3)
sens_rows = [
    ("1.05", "116.5", "971~1000", "25.3~43.0"),
    ("1.10", "122.1", "941~1000", "22.1~51.1"),
    ("1.25", "138.7", "855~1004", "16.1~60.1"),
    ("1.50", "166.5", "743~1004", "11.6~64.2"),
    ("2.00", "221.9", "562~1004", "8.7~70.8"),
]
three_line_table(["α", "门槛 J (m)", "离 S1 (m)", "偏角 (°)"], sens_rows)
para("")

# ================= 5 模型检验 =================
h1("5  模型检验")
h2("5.1  数值口径与误差分析")
para("本文 J 为有限源/误差网格上的最坏值估计：源网格加密至 17 方向 × 约 10 m 径向步长、误差网格加密至 17 点（步长 0.125°）后，三构型最坏值均收敛稳定；射线退出距离与独立解析二次方程求根（80 位十进制精度）对照的最大差异为 6.3×10⁻¹³ m；4.3 节的域裁剪数值实例进一步验证了后验定义的完备性。")
h2("5.2  灵敏度分析")
para("对容差系数 α 做 1.05~2.00 的梯度扰动（表 3）：α∈[1.1,1.5] 时候选区域形状稳定（贴外沿两瓣），结论不随容差取值漂移。对最优点做边界回退检验：中心构型最优点贴 C_recv 边界（裕量仅 0.02 m），向内回退约 3 m 至 (841,±544) 后 J 仅增至 111.3 m（+0.3%）、裕量增至 2 m——部署对位置误差不敏感。")
h2("5.3  稳健性检验")
para("其一，三构型一致：最坏情形均出现在楔形远弧端点 × 误差极值处，机理统一。其二，交叉验证：将本文候选点交由另一套独立实现（不同几何引擎、4096 源点、129 误差点、自适应加密）重评，三构型最坏直径绝对差分别为 0.0017 / 0.0005 / 0.0052 m（表 4），两套完全独立的数值管线互证到毫米级。其三，边界回归：近场强信号分支（后验取 Ω1∩B(q,5) 实际交集）、逐方向目标圆裁剪、重复地点、切向构型镜像点越界等边界情形的回归测试全部通过。")
caption("表 4  两套独立实现的交叉验证", before=8, after=3)
cv_rows = [
    ("中心", "111.2918", "111.293", "0.0017"),
    ("向外 800", "40.9714", "40.972", "0.0005"),
    ("切向", "8.1658", "8.171", "0.0052"),
]
three_line_table(["构型", "独立实现报告值 J (m)", "本文在其点重算 J (m)", "绝对差 (m)"], cv_rows)
para("")

# ================= 6 模型优缺点评价 =================
h1("6  模型优缺点评价")
h2("6.1  模型的优点")
adv = [
    "接收保证的推导严格：利用“S1 已收到信号 ⟹ R ≥ d1”把约束精化为 d2 ≤ max(1000, d1)，既避免“一律按 1000 m”的过度保守，也避免“一律按 1500 m”的假保证；三圆近似被明确定位为保守充分条件，保留其几何直观。",
    "后验区域完整裁剪：定位区域严格交回源可能域与接收上界，排除物理上不可能的源位置，避免了未裁剪口径约 20% 的直径虚高。",
    "最坏情形口径与任务一致：按有界误差集合取不利值、不依赖不可验证的分布假设，符合“确保清除”的稳健性要求。",
    "确定性、可复现：目标为一维标量且可几何精确计算，多级网格细化即可稳定求解，无需随机优化，结果逐位复现。",
    "数值口径诚实：如实标注 J 为有限网格上确界估计、连续全局最优未认证，不做虚假的连续证明。",
]
for i, s in enumerate(adv, 1):
    rich_para([(f"（{i}）", {"bold": True}), (s, {})])
para("")
h2("6.2  模型的缺点")
rich_para([("（1）", {"bold": True}), ("J 是有限源/误差网格上的最坏值估计，连续上确界与连续全局最优尚未严格认证，最优点附近（C_recv 边界）仍存在米级以内的数值不确定性。", {})])
rich_para([("（2）", {"bold": True}), ("对测向误差按有界最坏情形处理，结果偏保守；若任务提供误差分布信息，可用期望口径进一步优化。", {})])
rich_para([("（3）", {"bold": True}), ("未考虑 S2 的行走时间、能耗与地形约束，这些因素在问题三的多点序列问题中处理。", {})])
h2("6.3  模型的改进")
rich_para([("（1）", {"bold": True}), ("以区间算术（外向舍入）对 Ω1 与接收约束做严格认证，或借鉴分支定界[4] 的思想给出连续全局最优的间隙证明。", {})])
rich_para([("（2）", {"bold": True}), ("若获得误差分布信息，引入期望口径并与最坏口径对照，给出两种口径下的策略差异。", {})])
rich_para([("（3）", {"bold": True}), ("将本模型的 J 作为子目标嵌入问题三的多点序列规划，形成“粗定位→精细清除”的完整流程。", {})])

# ================= AI 工具使用声明 =================
h1("AI 工具使用声明")
para("本参赛队在竞赛过程中使用了 AI 工具（DeepSeek 等），主要用于：代码编写与调试、数值计算的独立交叉验证、公式推导核验与文档表述润色。论文中的模型、数据与结论均由本队推导、计算与验证，AI 工具的使用情况详见支撑材料，并在参考文献中列出所用工具。")

# ================= 参考文献 =================
h1("参考文献")
refs = [
    "[1] Storn R, Price K. Differential Evolution – A Simple and Efficient Heuristic for Global Optimization over Continuous Spaces[J]. Journal of Global Optimization, 1997, 11(4): 341-359.",
    "[2] Deb K, Pratap A, Agarwal S, et al. A fast and elitist multiobjective genetic algorithm: NSGA-II[J]. IEEE Transactions on Evolutionary Computation, 2002, 6(2): 182-197.",
    "[3] Metropolis N, Ulam S. The Monte Carlo Method[J]. Journal of the American Statistical Association, 1949, 44(247): 335-341.",
    "[4] Land A H, Doig A G. An Automatic Method of Solving Discrete Programming Problems[J]. Econometrica, 1960, 28(3): 497-520.",
    "[5] DeepSeek, DeepSeek-V3.2 / DeepSeek-R1 系列, 深度求索（DeepSeek）, 2026-09-12.",
]
for r_ in refs:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(r_)
    set_run(run, size=10.5)

# ================= 附录 =================
doc.add_page_break()
h1("附录")
h2("附录 1  支撑材料文件列表")
for f in ["code/problem2_model.py（问题二融合修正版模型）",
          "code/problem2_paper.py（三构型求解、出图与汇总）",
          "code/problem1_localize.py（问题一修正版几何引擎）",
          "code/test_problem2.py（14 项回归测试）",
          "code/final_refine.py（0.5 m 约束精搜）",
          "results/problem2_summary.json（全部数值结果与交叉验证）",
          "figs/problem2_candidate_region_*.png（三张候选区域图）"]:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run("· " + f)
    set_run(run, size=10.5)

h2("附录 2  核心模型源代码（problem2_model.py）")
with open(os.path.join(CODE, "problem2_model.py"), encoding="utf-8") as fh:
    code1 = fh.read()
for line in code1.splitlines():
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(line if line else " ")
    set_run(run, size=7.5, latin="Consolas", east="宋体")

doc.add_page_break()
h2("附录 3  求解与出图源代码（problem2_paper.py）")
with open(os.path.join(CODE, "problem2_paper.py"), encoding="utf-8") as fh:
    code2 = fh.read()
for line in code2.splitlines():
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(line if line else " ")
    set_run(run, size=7.5, latin="Consolas", east="宋体")

doc.add_page_break()
h2("附录 4  公式 LaTeX 源码（便于 Word 公式编辑器 Alt+= 录入）")
eq_tex = [
    ("式(1)", r"\Omega_1=\left\{g:\ \|g\|\le 1800,\ 5<\|g-S_1\|\le 1500,\ \left|\mathrm{wrap}\left(\theta(g-S_1)-\theta_1\right)\right|\le 1^\circ\right\}"),
    ("式(2)", r"R\ge \|G-S_1\|,\quad d_2=\|S_2-G\|\ \le\ L(G):=\max\left(1000,\ \|G-S_1\|\right)"),
    ("式(3)", r"C_{\mathrm{recv}}=\left\{q:\ F(q)=\sup_{g\in\Omega_1}\left(\|q-g\|-L(g)\right)\le 0\right\},\quad L(g)=\max(1000,\ \|g-S_1\|)"),
    ("式(4)", r"\Omega_2=W_1\cap W_2\cap\Omega_1\cap\left\{5<\|z-q\|\le 1500\right\}"),
    ("式(5)", r"J(q)=\sup_{g\in\Omega_1}\ \sup_{e_2\in[-1^\circ,1^\circ]}\ \mathrm{diam}\left(\Omega_2(q,g,e_2)\right)"),
    ("式(6)", r"\mathcal{R}_{\alpha}=\left\{q\in C_{\mathrm{recv}}:\ J(q)\le \alpha J_{\min}\right\}"),
    ("式(7)", r"D\approx \frac{2\tan 1^\circ}{\sin\psi}\sqrt{r_1^2+r_2^2+2r_1r_2|\cos\psi|}"),
    ("式(8)", r"C_{\mathrm{recv}}\supseteq B(S_1,1000)\cap B(P_-,1000)\cap B(P_+,1000)"),
]
for name, tex in eq_tex:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{name}：")
    set_run(r1, size=10.5, bold=True, east="宋体", latin="Times New Roman")
    r2 = p.add_run(tex)
    set_run(r2, size=10.5, latin="Consolas", east="宋体")

out = os.path.join(ROOT, "问题二论文.docx")
doc.save(out)
print("saved:", out)
